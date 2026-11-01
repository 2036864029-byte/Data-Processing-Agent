"""输出格式测试（T114，B7）——FR-018。

JSON / CSV / Markdown / Excel 四种交付格式；缺口页标记必须内嵌在产物中。
"""

import csv
import io
import json

from openpyxl import load_workbook

from dataagent.contracts.artifacts import PageArtifact
from dataagent.contracts.errors import DefectCode
from dataagent.pipeline.merge.merger import merge_document
from dataagent.pipeline.output.delivery import decide_delivery
from dataagent.pipeline.output.formats import to_csv, to_excel, to_json, to_markdown
from dataagent.pipeline.repair.fallback import FallbackPath
from dataagent.pipeline.repair.gap import PageGap

DOC_ID = "c" * 64


def artifact(page_no: int, markdown: str) -> PageArtifact:
    return PageArtifact(
        doc_id=DOC_ID,
        page_no=page_no,
        engine="engine.pdf@1.0.0",
        params_hash="params",
        markdown=markdown,
    )


def gap(page_no: int) -> PageGap:
    return PageGap(
        doc_id=DOC_ID,
        page_no=page_no,
        defect_codes=(DefectCode.TABLE_STRUCTURE_BROKEN,),
        fallback_used=FallbackPath.MARK_ONLY,
    )


def merged():
    return merge_document(
        DOC_ID,
        [artifact(1, "第一页\n"), artifact(3, "第三页\n")],
        [gap(2)],
    )


def test_json_contains_provenance_and_gap_flag() -> None:
    document = merged()
    payload = json.loads(to_json(document, decide_delivery(document)))

    assert payload["doc_id"] == DOC_ID
    assert payload["delivery_level"] == "needs_review"
    assert payload["pages"][1]["gap"] is True
    assert payload["pages"][1]["page_no"] == 2
    assert "dataagent:gap" in payload["pages"][1]["markdown"]
    assert payload["gap_pages"] == [2]


def test_markdown_keeps_page_order_and_gap_marker() -> None:
    document = merged()

    text = to_markdown(document, decide_delivery(document))

    assert text.index("第一页") < text.index("dataagent:gap")
    assert text.index("dataagent:gap") < text.index("第三页")


def test_csv_has_one_row_per_page_including_gaps() -> None:
    document = merged()

    rows = list(csv.DictReader(io.StringIO(to_csv(document, decide_delivery(document)))))

    assert [row["page_no"] for row in rows] == ["1", "2", "3"]
    assert [row["gap"] for row in rows] == ["false", "true", "false"]
    assert "table_structure_broken" in rows[1]["defect_codes"]


def test_excel_is_a_valid_workbook() -> None:
    document = merged()

    data = to_excel(document, decide_delivery(document))
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook.active

    assert sheet.max_row == len(document.pages) + 1
    assert sheet.cell(row=3, column=1).value == 2
    assert str(sheet.cell(row=3, column=2).value).lower() == "true"
